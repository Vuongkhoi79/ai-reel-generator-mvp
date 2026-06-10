import json
import os
import re
import threading
import urllib.error
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
DATABASE_PATH = Path(os.getenv("DATABASE_PATH", BASE_DIR / "DATABASE_MARKETING.xlsx"))
if not DATABASE_PATH.exists():
    DATABASE_PATH = BASE_DIR.parent / "DATABASE_MARKETING.xlsx"
HOST = "0.0.0.0"
PORT = int(os.environ.get("PORT", 8088))
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")
ZALO_URL = os.getenv("ZALO_URL", "https://zalo.me/g/9uomhrx1pwhmhosltoze")
GOOGLE_SHEET_WEBHOOK_URL = os.getenv("GOOGLE_SHEET_WEBHOOK_URL", "")
TRACKING_PATH = Path(os.getenv("TRACKING_PATH", BASE_DIR / "tracking_events.jsonl"))
TRACKING_LOCK = threading.Lock()


def normalize(text):
    return re.sub(r"\s+", " ", str(text or "").strip()).lower()


def load_database():
    wb = load_workbook(DATABASE_PATH, read_only=True, data_only=True)

    industries = []
    products = {}
    pains = {}
    goals = {}
    angles = {}

    for row in wb["Industries"].iter_rows(min_row=2, values_only=True):
        industry_id, industry, audience, positioning_note = row[:4]
        item = {
            "id": int(industry_id),
            "name": industry,
            "audience": audience,
            "positioningNote": positioning_note,
        }
        industries.append(item)
        products[item["id"]] = []
        pains[item["id"]] = []
        goals[item["id"]] = []
        angles[item["id"]] = []

    for row in wb["Products"].iter_rows(min_row=2, values_only=True):
        industry_id, industry, product_id, product_service, content_hint = row[:5]
        products[int(industry_id)].append({
            "id": int(product_id),
            "name": product_service,
            "hint": content_hint,
        })

    for row in wb["PainPoints"].iter_rows(min_row=2, values_only=True):
        industry_id, industry, pain_id, pain_point, reel_hook_hint = row[:5]
        pains[int(industry_id)].append({
            "id": int(pain_id),
            "name": pain_point,
            "hint": reel_hook_hint,
        })

    for row in wb["Goals"].iter_rows(min_row=2, values_only=True):
        industry_id, industry, goal_id, goal, cta_hint = row[:5]
        goals[int(industry_id)].append({
            "id": int(goal_id),
            "name": goal,
            "hint": cta_hint,
        })

    for row in wb["Angles"].iter_rows(min_row=2, values_only=True):
        industry_id, industry, angle_id, angle, angle_description, example_hook = row[:6]
        angles[int(industry_id)].append({
            "id": int(angle_id),
            "name": angle,
            "description": angle_description,
            "exampleHook": example_hook,
        })

    return {
        "industries": industries,
        "products": products,
        "pains": pains,
        "goals": goals,
        "angles": angles,
        "durations": ["15s", "25s", "45s", "60s"],
        "zaloUrl": ZALO_URL,
        "demoMode": not bool(os.getenv("OPENAI_API_KEY")),
    }


DATABASE = load_database()


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def item_name(items, item_id):
    try:
        return find_item(items, item_id)["name"]
    except Exception:
        return ""


def enrich_tracking_payload(payload):
    data = dict(payload or {})
    industry_id = data.get("industryId")
    if industry_id:
        try:
            industry_id = int(industry_id)
            data["industryName"] = item_name(DATABASE["industries"], industry_id)
            if data.get("productId"):
                data["productName"] = item_name(DATABASE["products"].get(industry_id, []), data["productId"])
            if data.get("painId"):
                data["painPoint"] = item_name(DATABASE["pains"].get(industry_id, []), data["painId"])
            if data.get("goalId"):
                data["goalName"] = item_name(DATABASE["goals"].get(industry_id, []), data["goalId"])
            if data.get("angleId"):
                data["angleName"] = item_name(DATABASE["angles"].get(industry_id, []), data["angleId"])
        except Exception:
            pass
    return data


def post_to_google_sheet(event):
    if not GOOGLE_SHEET_WEBHOOK_URL:
        return
    try:
        request = urllib.request.Request(
            GOOGLE_SHEET_WEBHOOK_URL,
            data=json.dumps(event, ensure_ascii=False).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        urllib.request.urlopen(request, timeout=8).read()
    except Exception as exc:
        print(f"Google Sheet tracking failed: {exc}")


def record_event(event_type, payload=None, handler=None):
    event = {
        "timestamp": now_iso(),
        "eventType": event_type,
        "payload": enrich_tracking_payload(payload or {}),
    }
    if handler:
        event["ip"] = handler.client_address[0] if handler.client_address else ""
        event["userAgent"] = handler.headers.get("User-Agent", "")

    TRACKING_PATH.parent.mkdir(parents=True, exist_ok=True)
    line = json.dumps(event, ensure_ascii=False)
    with TRACKING_LOCK:
        with TRACKING_PATH.open("a", encoding="utf-8") as file:
            file.write(line + "\n")
    post_to_google_sheet(event)
    return event


def read_tracking_events():
    if not TRACKING_PATH.exists():
        return []
    events = []
    with TRACKING_LOCK:
        with TRACKING_PATH.open("r", encoding="utf-8") as file:
            for line in file:
                line = line.strip()
                if not line:
                    continue
                try:
                    events.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    return events


def dashboard_metrics():
    events = read_tracking_events()
    visits = sum(1 for event in events if event.get("eventType") == "visit")
    generates = sum(1 for event in events if event.get("eventType") == "generate")
    zalo_clicks = sum(
        1
        for event in events
        if event.get("eventType") in ("zalo_click", "zalo_page_click", "zalo_group_click")
    )
    industry_counter = Counter(
        event.get("payload", {}).get("industryName") or "Không xác định"
        for event in events
        if event.get("eventType") == "generate"
    )
    conversion_rate = round((zalo_clicks / visits) * 100, 2) if visits else 0
    return {
        "visits": visits,
        "generates": generates,
        "zaloClicks": zalo_clicks,
        "conversionRate": conversion_rate,
        "topIndustries": [
            {"industry": industry, "count": count}
            for industry, count in industry_counter.most_common(10)
        ],
        "recentEvents": events[-50:][::-1],
        "googleSheetEnabled": bool(GOOGLE_SHEET_WEBHOOK_URL),
    }


def find_item(items, item_id):
    for item in items:
        if int(item["id"]) == int(item_id):
            return item
    raise ValueError("Không tìm thấy dữ liệu đã chọn.")


def build_prompt(payload):
    industry_id = int(payload["industryId"])
    industry = find_item(DATABASE["industries"], industry_id)
    product = find_item(DATABASE["products"][industry_id], payload["productId"])
    pain = find_item(DATABASE["pains"][industry_id], payload["painId"])
    goal = find_item(DATABASE["goals"][industry_id], payload["goalId"])
    angle = find_item(DATABASE["angles"][industry_id], payload["angleId"])
    duration = payload["duration"]

    return f"""Bạn là chuyên gia short-form content cho thị trường Việt Nam.

Hãy tạo nội dung Reel/TikTok/Shorts bằng tiếng Việt có dấu chuẩn, cụ thể, không chung chung.

INPUT:
- Ngành: {industry["name"]}
- Tệp khách hàng: {industry["audience"]}
- Sản phẩm/dịch vụ: {product["name"]}
- Pain point chính: {pain["name"]}
- Mục tiêu Reel: {goal["name"]}
- Angle content: {angle["name"]} - {angle["description"]}
- Thời lượng mỗi kịch bản: {duration}
- CTA ưu tiên: kéo inbox và kéo người vào Zalo Group nhận thêm mẫu

YÊU CẦU OUTPUT:
Trả về JSON hợp lệ, không markdown, không giải thích ngoài JSON.
Schema:
{{
  "hooks": ["5 hook 0-2 giây"],
  "scripts": [
    {{
      "title": "Tên kịch bản",
      "hook": "Hook chính của riêng kịch bản này",
      "duration": "{duration}",
      "scenes": [
        {{
          "time": "0-3s",
          "visual": "Cảnh quay/B-roll",
          "voice": "Lời thoại",
          "overlay": "Text trên màn hình"
        }}
      ],
      "cta": "CTA cuối video"
    }}
  ],
  "captions": ["10 caption đăng Facebook/TikTok"],
  "ctas": ["5 CTA khác nhau"],
  "imagePrompts": ["3 prompt tạo ảnh/thumbnail bằng Canva hoặc AI image"],
  "capcutGuide": ["Các bước dựng trong CapCut"]
}}

Quy tắc:
- Tạo đúng 5 hooks, 3 scripts, 10 captions, 5 ctas, 3 imagePrompts.
- Mỗi script phải có hook riêng, dùng được làm câu mở đầu video.
- Mỗi script phải đúng logic thời lượng {duration}; càng ngắn càng ít cảnh.
- Hook phải đánh thẳng vào pain point.
- Caption phải có CTA comment/inbox/vào Zalo nhưng không spam.
- CTA phải tự nhiên, phù hợp người bán hàng online.
- Prompt thumbnail phải mô tả bố cục, chủ thể, text overlay, màu sắc, tỷ lệ 9:16.
- Hướng dẫn CapCut phải thực dụng: template, text overlay, auto caption, voice, xuất video.
"""


def demo_response(payload):
    industry_id = int(payload["industryId"])
    industry = find_item(DATABASE["industries"], industry_id)
    product = find_item(DATABASE["products"][industry_id], payload["productId"])
    pain = find_item(DATABASE["pains"][industry_id], payload["painId"])
    goal = find_item(DATABASE["goals"][industry_id], payload["goalId"])
    angle = find_item(DATABASE["angles"][industry_id], payload["angleId"])
    duration = payload["duration"]

    product_name = product["name"]
    pain_name = pain["name"]
    industry_name = industry["name"]
    angle_name = angle["name"]
    goal_name = goal["name"]

    return {
        "hooks": [
            f"Bạn đang gặp {pain_name} mà vẫn làm theo cách này?",
            f"Sai lầm khiến {product_name} khó ra kết quả như mong đợi.",
            f"Nếu bạn thuộc nhóm khách của {industry_name}, xem hết {duration} này.",
            f"Đừng mua {product_name} trước khi biết điều này.",
            f"Một lỗi nhỏ có thể làm bạn mất tiền với {product_name}.",
        ],
        "scripts": [
            {
                "title": f"{angle_name}: lỗi phổ biến khi chọn {product_name}",
                "hook": f"90% người gặp {pain_name} đều đang mắc sai lầm này.",
                "duration": duration,
                "scenes": [
                    {"time": "0-3s", "visual": "Cận cảnh vấn đề hoặc biểu cảm bối rối", "voice": f"Nếu bạn đang {pain_name}, có thể bạn đang mắc lỗi này.", "overlay": "Sai lầm nhiều người mắc"},
                    {"time": "3-12s", "visual": "Minh họa cách làm sai", "voice": f"Nhiều người chọn {product_name} chỉ vì thấy quảng cáo hấp dẫn, nhưng không kiểm tra nhu cầu thật.", "overlay": "Chọn theo quảng cáo"},
                    {"time": "12-21s", "visual": "Hiển thị checklist 3 điểm", "voice": "Hãy kiểm tra vấn đề, mục tiêu và tiêu chí phù hợp trước khi quyết định.", "overlay": "Vấn đề - mục tiêu - tiêu chí"},
                    {"time": "21-25s", "visual": "Màn hình CTA/Zalo", "voice": "Muốn mình gửi checklist mẫu? Inbox hoặc vào Zalo nhận thêm.", "overlay": "Vào Zalo nhận mẫu"},
                ],
                "cta": "Inbox từ khóa AI hoặc vào Zalo Group để nhận thêm mẫu theo ngành.",
            },
            {
                "title": f"Trước khi dùng {product_name}, hãy tự hỏi 3 câu",
                "hook": f"Đừng chọn {product_name} nếu bạn chưa trả lời 3 câu này.",
                "duration": duration,
                "scenes": [
                    {"time": "0-4s", "visual": "Người dùng nhìn nhiều lựa chọn trên màn hình", "voice": f"Đừng vội chọn {product_name} nếu bạn chưa trả lời 3 câu này.", "overlay": "3 câu trước khi chọn"},
                    {"time": "4-14s", "visual": "Text từng câu hỏi xuất hiện", "voice": "Bạn đang cần giải quyết vấn đề gì? Kết quả mong muốn là gì? Ai sẽ hỗ trợ khi có lỗi?", "overlay": "Vấn đề / kết quả / hỗ trợ"},
                    {"time": "14-22s", "visual": "So sánh lựa chọn đúng và sai", "voice": f"Cách này giúp bạn tránh {pain_name} và ra quyết định rõ hơn.", "overlay": "Đỡ mất tiền, đỡ mất thời gian"},
                    {"time": "22-25s", "visual": "Nút inbox/Zalo", "voice": "Mình để mẫu checklist trong Zalo Group.", "overlay": "Nhận checklist miễn phí"},
                ],
                "cta": "Comment AI để nhận link Zalo và bộ mẫu miễn phí.",
            },
            {
                "title": f"Cách dùng {product_name} để {goal_name}",
                "hook": f"Muốn {goal_name}? Đừng bắt đầu bằng nội dung chung chung.",
                "duration": duration,
                "scenes": [
                    {"time": "0-3s", "visual": "Màn hình kết quả mong muốn", "voice": f"Muốn {goal_name}? Đừng bắt đầu bằng nội dung chung chung.", "overlay": "Đừng làm chung chung"},
                    {"time": "3-10s", "visual": "Chọn đúng pain point", "voice": f"Hãy mở đầu bằng vấn đề khách đang gặp: {pain_name}.", "overlay": "Bắt đầu từ pain"},
                    {"time": "10-19s", "visual": "Đưa giải pháp cụ thể", "voice": f"Sau đó nối sang {product_name} như một giải pháp cụ thể, dễ hiểu.", "overlay": "Nối pain với giải pháp"},
                    {"time": "19-25s", "visual": "CTA inbox/Zalo", "voice": "Cuối video chỉ cần một CTA: inbox hoặc vào Zalo nhận mẫu.", "overlay": "CTA rõ ràng"},
                ],
                "cta": "Vào Zalo Group nhận thêm 20 angle và caption mẫu.",
            },
        ],
        "captions": [
            f"Đang gặp {pain_name}? Lưu lại checklist này trước khi chọn {product_name}.",
            f"Một sai lầm nhỏ khi chọn {product_name} có thể làm bạn mất thời gian và tiền.",
            f"Muốn {goal_name}? Hãy bắt đầu từ pain thật của khách, không phải câu quảng cáo chung chung.",
            f"Đây là cách mình biến {angle_name.lower()} thành một Reel {duration}.",
            f"Nếu bạn bán trong ngành {industry_name}, mẫu này có thể dùng ngay hôm nay.",
            f"Comment AI nếu muốn nhận thêm mẫu hook và caption theo ngành.",
            f"Đừng để nội dung của bạn lặp lại mỗi ngày. Hãy đổi pain, đổi angle, đổi CTA.",
            f"Reel ngắn không cần phức tạp, chỉ cần đúng vấn đề và đúng lời kêu gọi hành động.",
            f"Mình để thêm mẫu trong Zalo Group cho ai cần làm content nhanh hơn.",
            f"Copy cấu trúc này: Pain thật -> lỗi phổ biến -> giải pháp -> CTA.",
        ],
        "ctas": [
            "Comment AI để nhận link Zalo Group.",
            "Inbox mình gửi mẫu theo ngành của bạn.",
            "Vào Zalo Group nhận thêm 20 hook miễn phí.",
            "Muốn mình tạo bản theo sản phẩm của bạn? Nhắn từ khóa REEL.",
            "Lưu video này và vào Zalo lấy checklist dựng CapCut.",
        ],
        "imagePrompts": [
            f"Ảnh thumbnail 9:16 cho {product_name}, chủ thể rõ, nền sáng sạch, text lớn: 'Đừng mắc lỗi này', phong cách hiện đại, dễ đọc trên điện thoại.",
            f"Thiết kế Canva 9:16 về pain point {pain_name}, bố cục trước/sau, màu tương phản, có khoảng trống cho tiêu đề ngắn.",
            f"Thumbnail Reel ngành {industry_name}, cảnh người dùng đang phân vân trước nhiều lựa chọn, overlay: '3 câu hỏi trước khi mua', phong cách thực tế Việt Nam.",
        ],
        "capcutGuide": [
            "Tạo project 9:16, chọn template đơn giản có nhịp cắt nhanh.",
            "Dán hook vào 2 giây đầu, dùng chữ lớn và tương phản.",
            "Chia script thành 3-4 đoạn, mỗi đoạn một text overlay ngắn.",
            "Bật Auto Captions, kiểm tra lại dấu tiếng Việt và từ chuyên ngành.",
            "Thêm B-roll sản phẩm/dịch vụ, giảm nhạc nền còn 10-20%, xuất 1080p.",
        ],
    }


def call_openai(prompt):
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("Missing OPENAI_API_KEY")

    body = {
        "model": OPENAI_MODEL,
        "input": [
            {"role": "system", "content": "Bạn chỉ trả về JSON hợp lệ theo schema người dùng yêu cầu."},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.8,
    }
    request = urllib.request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(body).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))

    text = data.get("output_text", "")
    if not text:
        chunks = []
        for item in data.get("output", []):
            for content in item.get("content", []):
                if content.get("type") == "output_text":
                    chunks.append(content.get("text", ""))
        text = "\n".join(chunks).strip()

    return json.loads(text)


def send_json(handler, status, data):
    payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(payload)))
    handler.end_headers()
    handler.wfile.write(payload)


class AppHandler(BaseHTTPRequestHandler):
    def log_message(self, format, *args):
        print(f"{self.address_string()} - {format % args}")

    def do_GET(self):
        path = self.path.split("?", 1)[0]
        if path in ("/api/data", "/api/options"):
            send_json(self, 200, DATABASE)
            return

        if path == "/api/dashboard":
            send_json(self, 200, dashboard_metrics())
            return

        if path == "/":
            path = "/index.html"

        file_path = (BASE_DIR / path.lstrip("/")).resolve()
        if not str(file_path).startswith(str(BASE_DIR)) or not file_path.exists():
            self.send_error(404)
            return

        content_types = {
            ".html": "text/html; charset=utf-8",
            ".css": "text/css; charset=utf-8",
            ".js": "application/javascript; charset=utf-8",
        }
        content = file_path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_types.get(file_path.suffix, "application/octet-stream"))
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def do_POST(self):
        if self.path == "/api/track":
            try:
                length = int(self.headers.get("Content-Length", "0"))
                payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else {}
                event_type = payload.pop("eventType", "custom")
                record_event(event_type, payload, self)
                send_json(self, 200, {"ok": True})
            except Exception as exc:
                send_json(self, 400, {"error": str(exc)})
            return

        if self.path != "/api/generate":
            self.send_error(404)
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            prompt = build_prompt(payload)
            if os.getenv("OPENAI_API_KEY"):
                try:
                    result = call_openai(prompt)
                    mode = "openai"
                except (urllib.error.URLError, TimeoutError, json.JSONDecodeError, RuntimeError) as exc:
                    result = demo_response(payload)
                    mode = f"demo_fallback: {exc}"
            else:
                result = demo_response(payload)
                mode = "demo"
            record_event("generate", {**payload, "mode": mode}, self)
            send_json(self, 200, {"mode": mode, "prompt": prompt, "result": result})
        except Exception as exc:
            send_json(self, 400, {"error": str(exc)})


if __name__ == "__main__":
    if not DATABASE_PATH.exists():
        raise SystemExit(f"Không tìm thấy database: {DATABASE_PATH}")
    server = ThreadingHTTPServer((HOST, PORT), AppHandler)
    print(f"Máy Tạo Reel AI đang chạy: http://{HOST}:{PORT}")
    print(f"Demo Mode: {'ON' if not os.getenv('OPENAI_API_KEY') else 'OFF'}")
    server.serve_forever()
